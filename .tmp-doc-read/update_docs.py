from pathlib import Path
from docx import Document


ROOT = Path(r"C:\documents\xiaoruan")


def iter_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def replace_everywhere(doc, replacements):
    for paragraph in iter_paragraphs(doc):
        text = paragraph.text
        updated = text
        for old, new in replacements:
            updated = updated.replace(old, new)
        if updated != text:
            paragraph.text = updated


def replace_paragraphs(doc, replacements):
    for paragraph in iter_paragraphs(doc):
        for needle, value in replacements:
            if needle in paragraph.text:
                paragraph.text = value


def append_section(doc, heading, paragraphs, bullets=()):
    doc.add_paragraph(heading, style="Heading 2")
    for text in paragraphs:
        doc.add_paragraph(text, style="Normal")
    for text in bullets:
        doc.add_paragraph(text, style="List Bullet")


def update_manual(path):
    doc = Document(path)
    replace_everywhere(
        doc,
        [
            ("知远", "晓软智能体"),
            ("OpenClaw Gateway", "统一 Agent 运行时"),
            ("OpenClaw 引擎", "Agent 运行时"),
            ("OpenClaw 会话", "Agent 会话"),
            ("OpenClaw 插件", "渠道连接器"),
            ("OpenClaw 工作目录", "应用工作区"),
            ("OpenClaw", "统一 Agent 运行时"),
            ("llama.cpp", "本地推理组件"),
        ],
    )
    replace_paragraphs(
        doc,
        [
            (
                "晓软智能体 是一款本地优先的桌面 AI Agent 应用",
                "晓软智能体是面向本地工作流的桌面 AI Agent 应用。当前定制版提供协作工作台、编程工作台、技能、MCP、专家、定时任务和消息渠道能力；模型调用必须经过用户绑定的 AISphere 平台。",
            ),
            (
                "晓软智能体 的设计目标是",
                "当前版本不提供产品账号登录、内置免费模型或第三方模型供应商入口。用户在设置中填写交付方提供的 AISphere 地址，应用只使用平台返回的可用模型；模型凭据由主进程管理。",
            ),
            (
                "本地推理 (Local Inference):",
                "模型来源: 通过 AISphere 绑定获取平台模型；当前定制版不开放第三方供应商、自定义模型或本地模型入口",
            ),
            (
                "多模型支持:",
                "模型目录: 由已绑定 AISphere 平台返回，平台下架或连接失败时不会自动替换模型",
            ),
            (
                "初始化 Agent 运行时引擎。",
                "初始化统一 Agent 运行时、模型网关和已启用的本地能力。",
            ),
            (
                "Agent 运行时启动过程中",
                "应用初始化和模型网关连接过程中",
            ),
            (
                "若引擎未就绪时尝试发送消息，系统将返回引擎未就绪的错误提示（错误代码: ENGINE_NOT_READY_CODE）。",
                "若模型网关尚未就绪，发送消息会提示连接或模型不可用；请先在设置中连接 AISphere 并选择平台模型。",
            ),
            (
                "AI 回复通过 SSE（Server-Sent Events）流式传输",
                "AI 回复通过应用内流式通道增量渲染",
            ),
            (
                "本地推理功能允许用户在本地计算机上运行 AI 模型",
                "当前定制版保留本地推理相关内部组件，但不向用户开放本地模型入口。内置模型请求仅使用已绑定 AISphere 平台返回的模型。",
            ),
            (
                "本地推理视图包含两个标签页：",
                "本地推理页面不属于当前定制版的用户模型配置流程；请在设置 → 模型中管理 AISphere 绑定。",
            ),
            (
                "本地推理子系统使用本地推理组件管理底层本地推理组件后端配置，支持 CPU 和 GPU 推理模式的切换。",
                "本地推理组件属于内部实现，不作为当前定制版的用户模型来源。",
            ),
            (
                "供应商配置：为每个供应商配置 API 密钥和 Base URL。",
                "AISphere 绑定：填写平台根地址，连接并获取平台返回的模型目录。",
            ),
            (
                "自定义供应商：支持最多 10 个自定义供应商。",
                "模型选择：只能选择当前 AISphere 平台返回的模型。",
            ),
            (
                "模型列表管理：查看、添加、编辑供应商的模型列表。",
                "模型目录：连接或刷新 AISphere 后查看可用模型；运行中的任务期间不能切换平台地址。",
            ),
            (
                "连接测试：测试供应商 API 连接状态。",
                "连接状态：显示 AISphere 是否可用；连接失败时清空可用模型并允许重试。",
            ),
            (
                "高级配置：支持设置超时时间、最大重试次数等参数。",
                "平台限制：平台地址仅支持 HTTP/HTTPS 根地址，模型请求通过应用内网关转发。",
            ),
            (
                "设置 → Agent Engine",
                "设置 → 模型",
            ),
            (
                "检查供应商 API 密钥是否填写正确。",
                "检查 AISphere 地址、网络连接和平台模型目录。",
            ),
            (
                "使用「连接测试」功能验证供应商 API 的连通性。",
                "点击「连接并获取模型」或「刷新模型」重试。",
            ),
            (
                "本地模型运行异常",
                "模型来源或平台连接异常",
            ),
            (
                "确认本地推理组件后端平台选择正确（CPU/GPU）。",
                "当前版本不提供本地模型入口；如平台模型不可用，请联系 AISphere 管理员。",
            ),
        ],
    )
    append_section(
        doc,
        "15. 编程工作台",
        [
            "编程工作台面向代码任务提供工作区、任务、泳道和 Agent 协作管理。它与普通 Cowork 会话分开保存，可在同一工作区中组织多个任务和执行泳道。",
            "内置编程 Agent 使用应用当前可用模型和权限策略；用户也可以发现并接入本机已安装且受信任的外部编程 Agent。外部 Agent 保留自己的账号、模型和配置，应用不会向其注入 AISphere 凭据。",
        ],
        [
            "创建或选择工作区，指定代码目录。",
            "创建任务并准备执行泳道，可配置 Agent、模型覆盖和可用命令。",
            "开始任务前查看权限请求；需要时暂停、取消、继续或恢复会话。",
            "查看工作区文件、Git 状态和差异，并执行暂存、提交、分支切换和拉取请求操作。",
            "通过交接和协作预设在不同 Agent 或泳道之间传递任务上下文。",
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def update_whitepaper(path):
    doc = Document(path)
    replace_everywhere(
        doc,
        [
            ("知远", "晓软智能体"),
            ("OpenClaw Gateway", "统一 Agent 运行时"),
            ("OpenClaw", "统一 Agent 运行时"),
            ("Pi SDK", "统一 Agent 运行时"),
            ("Pi", "统一 Agent 运行时"),
            ("llama.cpp", "本地推理组件"),
            ("API Key", "平台凭据"),
        ],
    )
    replace_paragraphs(
        doc,
        [
            (
                "晓软智能体 是一款面向开发者的桌面 AI Agent 平台",
                "晓软智能体是面向本地工作流的桌面 AI Agent 平台，采用 Electron + React 架构。当前定制版以 AISphere 作为唯一内置模型来源，并提供协作、编程、技能、MCP、定时任务和消息渠道能力。",
            ),
            (
                "可控性: 用户自主选择 AI 模型供应商，可完全使用本地模型，数据不出本地。",
                "可控性: 模型目录由用户绑定的 AISphere 平台提供，凭据由主进程管理；应用不提供产品账号、免费额度或第三方供应商入口。",
            ),
            (
                "进程隔离: 每个外部推理引擎",
                "进程隔离: Electron 主进程、渲染进程、渠道运行时和本地能力服务按边界隔离。",
            ),
            (
                "双内核架构",
                "统一 Agent 执行与渠道调度",
            ),
            (
                "配备两种 Agent 引擎实现，通过统一的 CoworkRuntime 接口进行抽象，按场景分工：",
                "Work、Chat、Channel 和 Cron 统一使用同一 Agent 执行内核；消息渠道和定时任务只负责接收、调度、投递和活动状态，不拥有独立的 Agent 任务状态。",
            ),
            (
                "OpenClaw 是 晓软智能体 的 Agent 运行时引擎",
                "Agent 运行时是晓软智能体的执行内核，负责会话、工具调用、权限请求和流式事件。",
            ),
            (
                "本地推理子系统以本地推理组件为核心引擎",
                "源码保留本地推理组件和模型管理能力，但当前定制版不向用户开放本地模型入口；内置请求只允许访问 AISphere 平台目录。",
            ),
            (
                "用户配置的 API 密钥存储在本地 SQLite 数据库。",
                "AISphere 地址和绑定状态按应用配置保存；实际平台凭据仅由主进程持有，不进入渲染进程配置或日志。",
            ),
            (
                "区域过滤机制，根据语言设置自动筛选可见的第三方平台。",
                "消息渠道按区域筛选：中文界面显示微信、钉钉、飞书、企业微信和 QQ；英文界面还显示 Telegram 和 Discord。",
            ),
            (
                "用户自行配置 API 密钥和本地模型。",
                "用户绑定 AISphere 地址并从平台目录选择模型；外部编程 Agent 的账号和模型配置保持在外部 Agent 自身。",
            ),
            (
                "高度可扩展: 技能体系、MCP 协议、自定义供应商、统一 Agent 运行时插件体系构成完整的扩展矩阵。",
                "高度可扩展: 技能、MCP、消息渠道、定时任务和编程 Agent 适配构成扩展矩阵。",
            ),
            (
                "本地优先: 核心功能不完全依赖云端，支持完全离线使用（本地模型模式）。",
                "本地优先: 会话、配置、技能和任务状态保存在本地；当前定制版模型调用依赖绑定的 AISphere，不承诺完全离线运行。",
            ),
        ],
    )
    append_section(
        doc,
        "15. 当前定制版边界",
        [
            "模型绑定：设置页只提供 AISphere 地址、连接/刷新和平台模型选择。平台发现接口和模型目录由主进程校验，模型请求通过本机网关转发。",
            "身份与更新：不提供产品账号登录、游客令牌、内置免费模型注册、第三方模型供应商配置或产品内自动更新入口。",
            "编程工作台：内置 Agent 与外部 ACP Agent 通过适配器隔离。外部 Agent 使用自己的账号、模型、工具和配置，不接收应用模型凭据。",
            "渠道与定时任务：渠道连接器承载外部消息和投递，定时任务负责触发和记录；两者均调用统一 Agent 执行内核。",
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def update_design(path):
    doc = Document(path)
    replace_everywhere(
        doc,
        [
            ("知远", "晓软智能体"),
            ("OpenClaw Gateway", "渠道与调度运行时"),
            ("OpenClaw", "渠道与调度运行时"),
            ("Pi SDK", "统一 Agent 执行内核"),
            ("Pi", "统一 Agent 执行内核"),
            ("llama.cpp", "本地推理组件"),
            ("rongxinai.db", "xiaoruan.sqlite"),
            ("OpenClaw 工作目录", "应用工作区"),
            ("openclaw_working_dir", "agent_workspace"),
        ],
    )
    replace_paragraphs(
        doc,
        [
            (
                "晓软智能体 是一款基于 Electron 框架构建的桌面端人工智能助手应用",
                "晓软智能体是基于 Electron 构建的桌面 AI Agent 应用。当前定制版的模型调用只经过用户绑定的 AISphere 平台；会话、技能、MCP、编程、定时任务和消息渠道状态由本地应用管理。",
            ),
            (
                "管理本地大语言模型的生命周期（下载、安装、启动、推理）。",
                "管理应用内可用的模型绑定和请求生命周期；本地推理相关代码不作为当前定制版的用户模型入口。",
            ),
            (
                "实现多渠道 IM 网关（企业微信、钉钉、飞书、微信、QQ等），使 Agent 可作为聊天机器人响应外部消息。",
                "通过渠道传输组件接入微信、钉钉、飞书、企业微信、QQ，以及英文界面的 Telegram 和 Discord，使 Agent 可响应外部消息。",
            ),
            (
                "OpenClawEngineManager         -- 统一 Agent 运行时引擎",
                "AISphereService              -- 模型平台绑定与目录管理",
            ),
            (
                "LlamaCppManager               -- 本地推理组件",
                "ChannelRuntime               -- 消息渠道与定时任务运行时",
            ),
            (
                "OpenClawEngineIpc: openclaw:engine:getStatus | openclaw:engine:install | ...",
                "ManagedProviderIpc: managed-provider:policy | managed-provider:catalog | ...",
            ),
            (
                "AuthIpc:         auth:login | auth:logout | auth:getTokens | ...",
                "CodingAgentIpc:  coding-agent:workspace | coding-agent:session | coding-agent:git | ...",
            ),
            (
                "OpenClawEngineManager         -- 统一 Agent 运行时引擎",
                "AISphereService              -- 模型平台绑定与目录管理",
            ),
            (
                "Agent 引擎路由",
                "统一 Agent 执行内核",
            ),
        ],
    )
    append_section(
        doc,
        "15. 当前定制版边界与编程工作台",
        [
            "模型配置由 AISphereService 统一管理。连接时先校验平台身份和模型目录，只收录运行中的平台模型；切换平台地址前必须停止运行中的任务。",
            "统一 Agent 执行内核处理 Work、Chat、Channel 和 Cron。渠道与定时任务组件只负责传输、触发、投递、并发控制和活动记录，不重复实现 Agent 会话状态。",
            "编程工作台通过 codingAgent 模块管理工作区、任务、泳道、外部 Agent Profile、权限、交接和 Git 操作。内置 Agent 使用应用模型配置；外部 Agent 通过 ACP 适配器运行，账号、模型和配置不跨边界注入。",
            "应用数据目录使用 XiaoruanAgent，SQLite 文件名为 xiaoruan.sqlite；旧产品目录和数据库不迁移、不读取。",
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


DOCS = sorted(ROOT.glob("*.docx"), key=lambda p: p.stat().st_size)
if len(DOCS) != 3:
    raise RuntimeError(f"expected three product documents, found {len(DOCS)}")

update_manual(DOCS[0])
update_whitepaper(DOCS[1])
update_design(DOCS[2])
