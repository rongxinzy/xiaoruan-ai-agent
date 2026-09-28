from pathlib import Path
from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


ROOT = Path(r"C:\documents\xiaoruan")


def iter_blocks(parent):
    if isinstance(parent, DocumentType):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise TypeError(type(parent))
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def cell_text(cell):
    parts = []
    for paragraph in cell.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in cell.tables:
        rows = [[cell_text(c) for c in row.cells] for row in table.rows]
        parts.append("; ".join(" | ".join(row) for row in rows))
    return "<br>".join(parts).replace("|", "\\|")


def table_markdown(table):
    rows = [[cell_text(cell) for cell in row.cells] for row in table.rows]
    if not rows:
        return []
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return lines


def paragraph_markdown(paragraph):
    text = paragraph.text.strip()
    if not text and paragraph._p.xpath(".//w:drawing"):
        text = "[图片]"
    if not text:
        return []
    style = paragraph.style.name if paragraph.style else "Normal"
    if style.startswith("Heading"):
        try:
            level = int(style.split()[-1])
        except ValueError:
            level = 2
        return ["#" * level + " " + text, ""]
    if style.startswith("List Bullet") or style.startswith("List Number"):
        return ["- " + text]
    if style == "CodeBlock" or "Code" in style:
        return ["@@CODE@@" + text]
    return [text, ""]


def convert(source, target):
    doc = Document(source)
    lines = []
    code_open = False

    def close_code():
        nonlocal code_open
        if code_open:
            lines.append("```")
            lines.append("")
            code_open = False

    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            rendered = paragraph_markdown(block)
            if rendered and rendered[0].startswith("@@CODE@@"):
                if not code_open:
                    lines.append("```")
                    code_open = True
                lines.append(rendered[0][8:])
                continue
            close_code()
            lines.extend(rendered)
        else:
            close_code()
            lines.extend(table_markdown(block))
            lines.append("")
    close_code()
    while lines and not lines[-1].strip():
        lines.pop()
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")


DOCS = [
    ("知远-2026.7.16-产品操作手册.docx", "知远-2026.7.16-产品操作手册.md"),
    ("知远-2026.7.16-产品技术白皮书.docx", "知远-2026.7.16-产品技术白皮书.md"),
    ("知远-2026.7.16-系统设计文档.docx", "知远-2026.7.16-系统设计文档.md"),
]

for source_name, target_name in DOCS:
    sources = [p for p in ROOT.glob("*.docx") if not p.name.startswith("~$") and p.stat().st_size]
    source = next((p for p in sources if p.name == source_name), None)
    if source is None:
        source = sorted(sources, key=lambda p: p.stat().st_size)[DOCS.index((source_name, target_name))]
    convert(source, ROOT / target_name)
