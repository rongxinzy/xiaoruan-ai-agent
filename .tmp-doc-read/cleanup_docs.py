from pathlib import Path
from docx import Document


ROOT = Path(r"C:\documents\xiaoruan")


def iter_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_all(doc, pairs):
    for paragraph in iter_paragraphs(doc):
        text = paragraph.text
        updated = text
        for old, new in pairs:
            updated = updated.replace(old, new)
        if updated != text:
            paragraph.text = updated


def set_matching(doc, pairs):
    for paragraph in iter_paragraphs(doc):
        for needle, value in pairs:
            if needle in paragraph.text:
                paragraph.text = value


def remove_row(row):
    row._tr.getparent().remove(row._tr)


def rewrite_manual_tables(doc):
    for table in doc.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:1] == ["供应商标识符"]:
            for row in list(table.rows)[1:]:
                remove_row(row)
            row = table.add_row()
            values = ["AISphere", "平台模型目录", "填写 AISphere 地址"]
            for cell, value in zip(row.cells, values):
                cell.text = value
        if header[:1] == ["项目"]:
            for row in table.rows:
                if row.cells and row.cells[0].text.strip() == "磁盘空间":
                    row.cells[1].text = "2 GB（应用本体）"
                    row.cells[2].text = "按工作区、技能和缓存需求准备"


def rewrite_whitepaper_tables(doc):
    for table in doc.tables:
        for row in table.rows:
            if not row.cells:
                continue
            first = row.cells[0].text.strip()
            if first == "模型供应商":
                row.cells[0].text = "模型来源"
                row.cells[1].text = "AISphere 平台目录（定制版独占）"
            elif first == "Agent 引擎":
                row.cells[0].text = "Agent 执行"
                row.cells[1].text = "统一 Agent 执行内核"
            elif first == "本地推理引擎":
                row.cells[0].text = "本地推理组件"
                row.cells[1].text = "源码保留，当前定制版不开放用户入口"
            elif first == "组件" and len(row.cells) > 1:
                if row.cells[1].text.strip() in {"统一 Agent 运行时", "本地推理组件（llama-server）"}:
                    row.cells[1].text = "应用运行时组件"
                    row.cells[2].text = "内部执行与调度，不作为产品配置入口"
            elif first == "依赖项":
                row.cells[0].text = "依赖边界"
                row.cells[1].text = "AISphere、Electron 内置运行环境和随包技能运行时"
                row.cells[2].text = "按交付环境提供"
            elif first == "类别":
                row.cells[0].text = "模型来源"
                row.cells[1].text = "AISphere 平台目录"


def cleanup_manual(path):
    doc = Document(path)
    replace_all(doc, [("API Key", "平台凭据"), ("本地模型", "平台模型")])
    set_matching(
        doc,
        [
            (
                "查看已安装的本地模型列表。",
                "当前版本不在用户界面提供本地模型列表；可用模型来自 AISphere 平台目录。",
            ),
            (
                "10 GB+ (含平台模型)",
                "按工作区、技能和缓存需求准备",
            ),
            (
                "自定义 Base URL + 平台凭据",
                "AISphere 地址",
            ),
        ],
    )
    rewrite_manual_tables(doc)
    doc.save(Path('.tmp-doc-read/clean-manual.docx'))


def cleanup_whitepaper(path):
    doc = Document(path)
    replace_all(
        doc,
        [
            ("openclaw-extensions/", "agent-extensions/"),
            ("openclawRuntimeAdapter.ts", "agentRuntimeAdapter.ts"),
            ("openclawEngineManager.ts", "channelRuntimeManager.ts"),
            ("openclawConfigSync.ts", "aisphere/service.ts"),
            ("openclawMemoryFile.ts", "memory/projectMemoryService.ts"),
            ("openclawSession", "channelSession"),
            ("openclawTokenProxy.ts", "aisphere/gateway.ts"),
            ("openclawAgentModels", "aisphereModelCatalog"),
            ("openclawLocalExtensions", "channelExtensions"),
            ("openclaw", "agent-runtime"),
            ("OpenClaw", "Agent 运行时"),
            ("llamacpp", "localInference"),
        ],
    )
    set_matching(
        doc,
        [
            (
                "本地扩展: 支持加载 agent-extensions/ 目录下的本地扩展模块。",
                "扩展能力: 通过技能、MCP 和编程 Agent 适配器扩展能力。",
            ),
            (
                "本地模型推理延迟取决于硬件配置（GPU 型号、显存大小）。",
                "模型响应延迟主要取决于 AISphere 平台和网络条件；本地推理组件不属于当前定制版用户入口。",
            ),
            (
                "运行本地模型时额外内存占用取决于模型大小（如 7B 参数量化模型约需 4-8 GB）。",
                "应用内存占用取决于会话、技能、附件和工作区规模；平台模型推理资源由 AISphere 服务承担。",
            ),
            (
                "8.3 自定义供应商",
                "8.3 模型绑定边界",
            ),
            (
                "10 个可自定义供应商",
                "仅允许绑定一个 AISphere 平台地址",
            ),
            (
                "本地模型推理",
                "本地推理组件（当前定制版不开放）",
            ),
        ],
    )
    rewrite_whitepaper_tables(doc)
    doc.save(Path('.tmp-doc-read/clean-whitepaper.docx'))


def cleanup_design(path):
    doc = Document(path)
    replace_all(
        doc,
        [
            ("openclawRuntimeAdapter.ts", "agentRuntimeAdapter.ts"),
            ("openclawEngineManager.ts", "channelRuntimeManager.ts"),
            ("openclawConfigSync.ts", "aisphere/service.ts"),
            ("openclawMemoryFile.ts", "memory/projectMemoryService.ts"),
            ("openclawSessionPolicy", "channelSessionPolicy"),
            ("openclawSession", "channelSession"),
            ("openclawTokenProxy.ts", "aisphere/gateway.ts"),
            ("openclawAgentModels", "aisphereModelCatalog"),
            ("openclawLocalExtensions", "channelExtensions"),
            ("openclaw", "agent-runtime"),
            ("OpenClaw", "Agent 运行时"),
            ("llamacpp", "localInference"),
        ],
    )
    set_matching(
        doc,
        [
            (
                "7.1 双内核架构",
                "7.1 统一 Agent 执行内核",
            ),
            (
                "系统采用双内核（Dual-Kernel）设计，两个引擎按场景分工，通过统一的 CoworkRuntime 接口向渲染进程屏蔽差异：",
                "Work、Chat、Channel 和 Cron 统一使用同一 Agent 执行内核。渠道与定时任务只负责接收、调度、投递和活动状态，不拥有独立的 Agent 任务状态。",
            ),
            (
                "export type CoworkAgentEngine = 'agent-runtime' | 'pi';",
                "Agent 执行内核由统一适配层提供，具体实现不作为产品协议暴露。",
            ),
            (
                "系统采用双内核",
                "系统采用统一执行内核",
            ),
            (
                "Agent 运行时 Runtime Adapter (agentRuntimeAdapter.ts)",
                "渠道与定时任务适配层",
            ),
            (
                "本地模型扫描与目录管理",
                "内部本地推理资源管理（当前定制版不开放用户入口）",
            ),
            (
                "模型与 Agent 运行时绑定",
                "AISphere 模型绑定",
            ),
            (
                "模型可作为 Agent 运行时的 Provider 接入",
                "当前定制版不把本地模型注册为用户可选 Provider；内置模型仅来自 AISphere 目录。",
            ),
            (
                "令牌代理通过 agent-runtimeTokenProxy.ts 实现，管理 API 密钥的安全传递。",
                "AISphere 网关负责校验请求范围并在主进程内转发平台凭据。",
            ),
            (
                "Agent 运行时 运行时构建 (agent-runtime:runtime:*): 从 GitHub 仓库克隆、编译 Agent 运行时。",
                "渠道运行时构建：准备随包的消息渠道传输组件。",
            ),
        ],
    )
    doc.save(Path('.tmp-doc-read/clean-design.docx'))


DOCS = sorted(
    (p for p in ROOT.glob("*.docx") if not p.name.startswith("~$")),
    key=lambda p: p.stat().st_size,
)
if len(DOCS) != 3:
    raise RuntimeError(f"expected three product documents, found {len(DOCS)}")
cleanup_manual(DOCS[0])
cleanup_whitepaper(DOCS[1])
cleanup_design(DOCS[2])
